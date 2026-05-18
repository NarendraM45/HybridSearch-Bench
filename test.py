import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_LINE_SPACING

# ===== CONFIG =====
SOURCE_FOLDER = r"C:\Users\naren\Python\HybridSearch Bench"
OUTPUT_DOCX = "project_code_dump.docx"

# Folders to exclude
EXCLUDED_FOLDERS = {
    "__pycache__",
    ".git",
    "data"
}

# File extensions to include
INCLUDE_EXTENSIONS = {
    ".py", ".txt", ".md", ".json", ".yaml", ".yml",
    ".toml", ".env", ".ini", ".cfg", ".dockerignore",
    ".gitignore", ".xml", ".html", ".css", ".js", ".env", ".example",
    ".ts", ".sql", ".sh", ".bat"
}


def is_excluded(path_parts):
    return any(part in EXCLUDED_FOLDERS for part in path_parts)


def should_include(file_path):
    filename = os.path.basename(file_path)

    # include extensionless important files
    extensionless = {
        "Dockerfile",
        "Makefile",
        ".env",
        ".env.example"
    }

    ext = os.path.splitext(filename)[1].lower()

    return (
        filename in extensionless
        or ext in INCLUDE_EXTENSIONS
    )


def add_file_to_doc(doc, root_folder, file_path):
    relative_path = os.path.relpath(file_path, root_folder)

    # compact file separator
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    run = p.add_run(f"\n===== FILE: {relative_path} =====\n")
    run.bold = True
    run.font.size = Pt(7)

    try:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()

    except UnicodeDecodeError:
        try:
            with open(file_path, "r", encoding="latin-1") as f:
                content = f.read()
        except Exception as e:
            content = f"[ERROR READING FILE: {e}]"

    except Exception as e:
        content = f"[ERROR READING FILE: {e}]"

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1

    run = p.add_run(content)
    run.font.name = "Consolas"
    run.font.size = Pt(6.5)  # smallest readable font


def main():
    doc = Document()

    # ultra compact margins
    section = doc.sections[0]
    section.top_margin = Pt(10)
    section.bottom_margin = Pt(10)
    section.left_margin = Pt(10)
    section.right_margin = Pt(10)

    files_added = 0

    for root, dirs, files in os.walk(SOURCE_FOLDER):

        # remove excluded folders from traversal
        dirs[:] = [
            d for d in dirs
            if d not in EXCLUDED_FOLDERS
        ]

        path_parts = root.split(os.sep)

        if is_excluded(path_parts):
            continue

        for file in sorted(files):
            file_path = os.path.join(root, file)

            if should_include(file_path):
                print(f"Adding: {file_path}")
                add_file_to_doc(
                    doc,
                    SOURCE_FOLDER,
                    file_path
                )
                files_added += 1

    doc.save(OUTPUT_DOCX)

    print("\nDone!")
    print(f"Files added: {files_added}")
    print(f"Saved as: {OUTPUT_DOCX}")


if __name__ == "__main__":
    main()